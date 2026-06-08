from __future__ import annotations

import math
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import Activity, RawActivity


REQUIRED_COLUMNS = (
    "Začetek",
    "Zaključek",
    "Kategorija",
    "Naziv",
    "Opis",
    "Sodelujoči",
    "Vpisal",
)

CATEGORY_MAPPING = {
    "delo v gasilnem domu": "Delo v domu",
    "delo v pisarni": "Delo v domu",
    "dežurstvo": "Dežurstva",
    "drugo": "Drugo",
    "obiski": "Drugo",
    "preventiva": "Drugo",
    "občni zbor": "Drugo",
    "dejavnosti na nivoju gz/gpo": "Drugo",
    "druženje": "Drugo",
    "izposoja gasilske opreme": "Drugo",
    "izposoja miz in klopi": "Drugo",
    "izposoja inventarja": "Drugo",
    "redarstvo": "Drugo",
    "pogreb": "Gasilska žalovanja",
    "žalna seja": "Gasilska žalovanja",
    "parada": "Gasilske prireditve",
    "požarna straža": "Požarna straža",
    "pregled vozil in opreme": "Pregled/servisiranje opreme",
    "pregledi hidrantnega omrežja": "Pregledi hidrantnega omrežja",
    "prevozi vode": "Prevozi vode",
    "seja nadzornega odbora": "Seja/sestanek",
    "seja upravnega odbora": "Seja/sestanek",
    "sestanek": "Seja/sestanek",
    "posvet/seminar": "Posvet / seminar",
    "urejanje okolice": "Urejanje okolice",
    "usposabljanje/izobraževanje": "Usposabljanje/izobraževanje",
    "vaje mladina": "Vaje",
    "vaje člani": "Vaje",
    "vaje veterani": "Vaje",
    "kondicijska vožnja": "Vaje",
}


class WorkbookError(ValueError):
    pass


def read_activities(path: str | Path) -> list[RawActivity]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    header_values = [cell.value for cell in next(worksheet.iter_rows(min_row=2, max_row=2))]
    headers = {str(value).strip(): index for index, value in enumerate(header_values) if value}
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise WorkbookError(f"Missing required columns: {', '.join(missing)}")

    activities: list[RawActivity] = []
    for excel_row_number, row in enumerate(
        worksheet.iter_rows(min_row=3, values_only=True),
        start=3,
    ):
        if not any(row):
            continue
        get = lambda name: _cell_text(row[headers[name]])
        try:
            start = parse_datetime(row[headers["Začetek"]])
            end = parse_datetime(row[headers["Zaključek"]])
        except ValueError as exc:
            raise WorkbookError(f"Row {excel_row_number}: {exc}") from exc

        activities.append(
            RawActivity(
                row_number=excel_row_number,
                start=start,
                end=end,
                category=get("Kategorija"),
                title=get("Naziv"),
                description=get("Opis"),
                participants_text=get("Sodelujoči"),
                entered_by=get("Vpisal"),
            )
        )

    return activities


def transform_activity(raw: RawActivity) -> Activity:
    warnings: list[str] = []
    vulkan_category = translate_category(raw.category)
    if vulkan_category.startswith("Neznana"):
        warnings.append(f"Unknown Ignis category: {raw.category}")

    hours = rounded_half_hours(raw.start, raw.end)
    if hours <= 0:
        warnings.append("Duration is zero or negative.")

    return Activity(
        row_number=raw.row_number,
        start=raw.start,
        end=raw.end,
        ignis_category=raw.category,
        vulkan_category=vulkan_category,
        title=raw.title,
        vulkan_note=raw.title,
        description=raw.description,
        participants=tuple(parse_participants(raw.participants_text)),
        entered_by=raw.entered_by,
        hours_value=hours,
        hours_text=format_hours(hours),
        date_text=format_vulkan_date(raw.start),
        warnings=tuple(warnings),
    )


def activities_for_month(raw_activities: Iterable[RawActivity], year: int, month: int) -> list[Activity]:
    selected = [
        transform_activity(raw)
        for raw in raw_activities
        if raw.start.year == year and raw.start.month == month
    ]
    return sorted(selected, key=lambda activity: activity.start)


def translate_category(ignis_category: str) -> str:
    clean = ignis_category.strip().lower()
    if not clean:
        return "Brez IGNIS vnosa"
    return CATEGORY_MAPPING.get(clean, "Neznana kategorija v VULKAN-u")


def rounded_half_hours(start: datetime, end: datetime) -> float:
    minutes = (end - start).total_seconds() / 60
    if minutes <= 0:
        return 0.0
    rounded = round((minutes / 60) * 2) / 2
    return max(0.5, rounded)


def format_hours(hours: float) -> str:
    if math.isclose(hours, round(hours)):
        return str(int(round(hours)))
    return f"{hours:.1f}".replace(".", ",")


def format_vulkan_date(date_value: datetime) -> str:
    return f"{date_value.day:02d}{date_value.month}"


def parse_participants(raw_text: str) -> list[str]:
    if not raw_text.strip():
        return []

    text = re.sub(r"\bVodja\s*:\s*", "", raw_text, flags=re.IGNORECASE)
    text = re.sub(r"\bSodelujoči\s*:\s*", ",", text, flags=re.IGNORECASE)
    parts = [part.strip() for part in re.split(r",|\n|;", text) if part.strip()]

    seen: set[str] = set()
    people: list[str] = []
    for part in parts:
        normalized = normalize_name(part)
        if normalized and normalized not in seen:
            seen.add(normalized)
            people.append(part)
    return people


def normalize_name(name: str) -> str:
    return " ".join(name.casefold().replace("\xa0", " ").split())


def parse_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if value is None:
        raise ValueError("missing datetime")

    text = str(value).strip()
    for fmt in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H.%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"invalid datetime '{text}'")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
